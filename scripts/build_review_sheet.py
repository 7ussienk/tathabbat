#!/usr/bin/env python3
"""يولّد review/review_sheet.xlsx للمرشد الشرعي من data/curated/widespread.jsonl.

سكربت بيانات (ليس كود منتج). يحتاج: pip install openpyxl
التشغيل:  python scripts/build_review_sheet.py
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
CURATED = ROOT / "data" / "curated" / "widespread.jsonl"
EXCLUDED = ROOT / "data" / "curated" / "excluded_from_demo.json"
MANIFEST = ROOT / "data" / "sources" / "manifest.json"
OUT = ROOT / "review" / "review_sheet.xlsx"

VERDICT_AR = {
    "authentic": "صحيح (authentic)",
    "weak": "ضعيف (weak)",
    "fabricated": "موضوع/مختلَق (fabricated)",
    "no_basis_per_scholar": "لا أصل له عند عالم — حكم منقول (no_basis_per_scholar)",
    "not_found_in_sources": "لم نجده في المصادر المفحوصة — امتناع (not_found_in_sources)",
    "disputed": "خلافي (disputed)",
    "misattributed": "قول لغير النبي ﷺ نُسب إليه (misattributed)",
    "quran_verified": "آية صحيحة النص (quran_verified)",
    "quran_misquoted": "آية منقولة بخطأ (quran_misquoted)",
}

HIGH_PRIORITY = {"W001", "W002", "W003", "W004", "W005", "W006", "W007", "W008", "W023", "W024", "W025", "W026"}
SHEET_ORDER = ["خلافية", "منتشرة", "صحيحة", "آيات"]  # الأولوية العالية أولاً

HEADERS = [
    "المعرّف",
    "أولوية",
    "النص المتداول",
    "الحكم المقترح",
    "الحكم المنقول بنصه",
    "المصدر والموضع",
    "الرابط",
    "ملاحظاتك",
    "موافق؟ (نعم/لا/تعديل)",
    "ملاحظة المراجع",
]
WIDTHS = [10, 10, 40, 26, 60, 38, 38, 50, 18, 40]
REVIEWER_FROM = len(HEADERS) - 1  # آخر عمودين للمراجع (1-based: 9 و10)

HEAD_FILL = PatternFill("solid", fgColor="1F4E5F")
REVIEWER_FILL = PatternFill("solid", fgColor="FFF2CC")
REVIEWER_HEAD_FILL = PatternFill("solid", fgColor="BF8F00")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)


def load_jsonl(path):
    return [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]


def main():
    for _s in (sys.stdout, sys.stderr):
        if hasattr(_s, "reconfigure"):
            _s.reconfigure(encoding="utf-8")
    manifest = {m["id"]: m for m in json.loads(MANIFEST.read_text(encoding="utf-8"))}
    excluded = {e["id"] for e in json.loads(EXCLUDED.read_text(encoding="utf-8"))["excluded_from_demo"]} if EXCLUDED.exists() else set()
    entries = load_jsonl(CURATED)

    def row_for(e):
        quotes, places, urls = [], [], []
        for s in e["sources"]:
            q = s.get("grading_quote") or s.get("quoted_text") or ""
            if q:
                if s.get("attribution_note"):
                    q += f"\n[{s['attribution_note']}]"
                quotes.append(q)
            title = manifest.get(s["source_id"], {}).get("title", s["source_id"])
            places.append(f"{title} — {s['location']}")
            urls.append(s.get("url") or "data/quran/quran-uthmani.txt (Tanzil)")
        notes = []
        if e.get("note"):
            notes.append(e["note"])
        if e.get("related_narration"):
            rn = e["related_narration"]
            notes.append(f"رواية متصلة: {rn['text']} — {rn.get('grading_quote', '')}")
        if e.get("authentic_alternative"):
            aa = e["authentic_alternative"]
            notes.append(f"البديل الصحيح المقترح: {aa['text']} ({aa['source_id']} {aa['location']})")
        if e["id"] in excluded:
            notes.append("مستبعد من العرض (excluded_from_demo)؛ يبقى في التقييم")
        return [
            e["id"],
            "عالية" if e["id"] in HIGH_PRIORITY else "عادية",
            e["claim_text"],
            VERDICT_AR.get(e["verdict"], e["verdict"]),
            "\n\n".join(quotes),
            "\n".join(places),
            "\n".join(urls),
            "\n".join(notes),
            None,
            None,
        ]

    def prio_sorted(items):  # عالية أولاً ثم بحسب المعرّف
        return sorted(items, key=lambda e: (e["id"] not in HIGH_PRIORITY, e["id"]))

    groups_raw = {
        "منتشرة": [e for e in entries if e["id"].startswith("W") and e["verdict"] != "disputed"],
        "خلافية": [e for e in entries if e["verdict"] == "disputed"],
        "صحيحة": [e for e in entries if e["id"].startswith("A")],
        "آيات": [e for e in entries if e["id"].startswith("Q")],
    }
    groups = {k: prio_sorted(groups_raw[k]) for k in SHEET_ORDER}

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "تعليمات"
    ws0.sheet_view.rightToLeft = True
    ws0.column_dimensions["A"].width = 22
    ws0.column_dimensions["B"].width = 110
    ws0["A1"] = "المطلوب منك"
    ws0["A1"].font = Font(bold=True, size=14)
    ws0["B1"] = ("راجع كل ورقة (منتشرة، خلافية، صحيحة، آيات): تحقق أن النص المنقول مطابق للمصدر وأن الحكم المقترح يوافق ما نُقل، "
                 "ثم اختر في عمود «موافق؟» نعم أو لا أو تعديل.")
    ws0["B2"] = ("اكتب في «ملاحظة المراجع» سبب الرفض أو التعديل المطلوب. العمودان الملوّنان لك وحدك؛ "
                 "لا تعدّل بقية الأعمدة. الأحكام منقولة بنصها من كتب القائمة البيضاء، والأداة ذكاء اصطناعي وليست مفتياً.")
    ws0["B3"] = ("التحقق الآلي من مطابقة النصوص للمصدر تم؛ المطلوب منك الحكم الشرعي على صحة العزو والتصنيف. "
                 "المدخلات ذات الأولوية «عالية» تأتي أولاً في كل ورقة، والأوراق مرتبة بحيث تأتي الأولوية العالية أولاً.")
    ws0["A4"] = "اسم المراجع:"
    ws0["A4"].font = Font(bold=True)
    ws0["B4"].fill = REVIEWER_FILL
    ws0["B4"].border = BORDER
    for r in (1, 2, 3):
        ws0[f"B{r}"].alignment = WRAP
        ws0.row_dimensions[r].height = 48

    dv_text = '"نعم,لا,تعديل"'
    for name, items in groups.items():
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(HEADERS)).column_letter}{max(len(items) + 1, 2)}"
        for c, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = REVIEWER_HEAD_FILL if c >= REVIEWER_FROM else HEAD_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center", readingOrder=2)
            cell.border = BORDER
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
        ws.row_dimensions[1].height = 32
        for r, e in enumerate(items, start=2):
            for c, v in enumerate(row_for(e), start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment = WRAP
                cell.border = BORDER
                if c >= REVIEWER_FROM:
                    cell.fill = REVIEWER_FILL
                if c == 2 and v == "عالية":
                    cell.fill = PatternFill("solid", fgColor="F8CBAD")
                    cell.font = Font(bold=True)
        dv = DataValidation(type="list", formula1=dv_text, allow_blank=True)
        dv.error = "اختر: نعم أو لا أو تعديل"
        ws.add_data_validation(dv)
        if items:
            col = ws.cell(row=1, column=REVIEWER_FROM).column_letter
            dv.add(f"{col}2:{col}{len(items) + 1}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"saved {OUT}  " + "، ".join(f"{k}: {len(v)}" for k, v in groups.items()))


if __name__ == "__main__":
    sys.exit(main())
