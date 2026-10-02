#!/usr/bin/env python3
"""يقرأ review/review_sheet.xlsx بعد مراجعة المرشد ويحدّث reviewed في البيانات.

سكربت بيانات (ليس كود منتج). يحتاج: pip install openpyxl
التشغيل:
  python scripts/apply_review.py --dry-run          # معاينة دون كتابة
  python scripts/apply_review.py                    # تطبيق
  python scripts/apply_review.py --reviewer "الاسم"  # يتجاوز الاسم المكتوب في ورقة التعليمات

القواعد:
  - «نعم»   => reviewed:true مع reviewed_by وreviewed_on، وإزالة أي review_status قديم.
  - «لا» / «تعديل» => reviewed:false مع review_status (rejected | needs_edit) وreview_note.
  - فارغ    => لا تغيير.
  - حالات eval/dataset.jsonl التي لها curated_ref: تصبح reviewed:true إذا وافق المراجع على **كل** مراجعها،
    وتبقى false (مع review_status) إذا رُفض أو طُلب تعديل أحدها. الحالات بلا curated_ref
    (فتوى، اصطناعية...) لا تغطيها هذه الورقة وتُراجع على حدة.
  - يُلحق صفاً في docs/REVIEW_LOG.md.
"""
import argparse
import datetime
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
H_ID, H_OK, H_NOTE = "المعرّف", "موافق؟ (نعم/لا/تعديل)", "ملاحظة المراجع"
STATUS = {"نعم": "approved", "لا": "rejected", "تعديل": "needs_edit"}


def dumps(o):
    return json.dumps(o, ensure_ascii=False, separators=(", ", ": "))


def read_jsonl(p):
    return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]


def write_jsonl(p, rows):
    p.write_text("\n".join(dumps(r) for r in rows) + "\n", encoding="utf-8", newline="\n")


def main():
    for _s in (sys.stdout, sys.stderr):
        if hasattr(_s, "reconfigure"):
            _s.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=str(ROOT), help="جذر المستودع (للاختبار على نسخة)")
    ap.add_argument("--xlsx", default=None)
    ap.add_argument("--reviewer", default=None)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    root = Path(a.root)
    xlsx = Path(a.xlsx) if a.xlsx else root / "review" / "review_sheet.xlsx"
    curated_p = root / "data" / "curated" / "widespread.jsonl"
    dataset_p = root / "eval" / "dataset.jsonl"
    log_p = root / "docs" / "REVIEW_LOG.md"
    today = datetime.date.today().isoformat()

    wb = load_workbook(xlsx, data_only=True)
    reviewer = (a.reviewer or "").strip()
    if not reviewer:
        reviewer = str(wb[wb.sheetnames[0]]["B4"].value or "").strip()
    if not reviewer:
        sys.exit("اسم المراجع مطلوب: اكتبه في ورقة التعليمات (B4) أو مرّر --reviewer")

    decisions = {}  # id -> (status, note)
    errors = []
    for ws in wb.worksheets[1:]:
        head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        try:
            i_id, i_ok, i_note = head.index(H_ID), head.index(H_OK), head.index(H_NOTE)
        except ValueError:
            errors.append(f"الورقة «{ws.title}»: ترويسة الأعمدة غير مطابقة")
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            rid = str(row[i_id] or "").strip()
            ok = str(row[i_ok] or "").strip()
            if not rid or not ok:
                continue
            if ok not in STATUS:
                errors.append(f"{rid}: قيمة غير مقبولة «{ok}» (المقبول: نعم/لا/تعديل)")
                continue
            note = str(row[i_note] or "").strip()
            if STATUS[ok] != "approved" and not note:
                errors.append(f"{rid}: «{ok}» بلا ملاحظة مراجع")
            decisions[rid] = (STATUS[ok], note)

    curated = read_jsonl(curated_p)
    ids = {e["id"] for e in curated}
    for rid in decisions:
        if rid not in ids:
            errors.append(f"{rid}: غير موجود في widespread.jsonl")
    if errors:
        print("أخطاء في ورقة المراجعة، لم يُكتب شيء:")
        for e in errors:
            print(" -", e)
        sys.exit(1)

    for e in curated:
        d = decisions.get(e["id"])
        if not d:
            continue
        status, note = d
        for k in ("review_status", "review_note", "reviewed_by", "reviewed_on"):
            e.pop(k, None)
        if status == "approved":
            e["reviewed"] = True
            e["reviewed_by"] = reviewer
            e["reviewed_on"] = today
            if note:
                e["review_note"] = note
        else:
            e["reviewed"] = False
            e["review_status"] = "rejected" if status == "rejected" else "needs_edit"
            e["review_note"] = note

    dataset = read_jsonl(dataset_p)
    ds_changed = 0
    for r in dataset:
        refs = [x["curated_ref"] for x in r["expected"] if x.get("curated_ref")]
        if not refs or not any(x in decisions for x in refs):
            continue
        sts = [decisions.get(x, ("pending", ""))[0] for x in refs]
        if all(s == "approved" for s in sts):
            new = (True, None)
        elif any(s in ("rejected", "needs_edit") for s in sts):
            new = (False, "needs_edit")
        else:
            continue
        r.pop("review_status", None)
        r["reviewed"] = new[0]
        if new[1]:
            r["review_status"] = new[1]
        ds_changed += 1

    n = {s: sum(1 for v in decisions.values() if v[0] == s) for s in ("approved", "rejected", "needs_edit")}
    print(f"المراجع: {reviewer} | موافق: {n['approved']} | رفض: {n['rejected']} | تعديل: {n['needs_edit']} | حالات eval متأثرة: {ds_changed}")
    if a.dry_run:
        print("(معاينة فقط — لم يُكتب شيء)")
        return
    write_jsonl(curated_p, curated)
    write_jsonl(dataset_p, dataset)
    row = f"| {today} | {reviewer} | {len(decisions)} مدخلاً (موافق {n['approved']}، رفض {n['rejected']}، تعديل {n['needs_edit']}) | من review/review_sheet.xlsx؛ حالات eval بلا curated_ref (فتوى/اصطناعية) لم تُراجع بهذه الورقة |\n"
    text = log_p.read_text(encoding="utf-8").rstrip("\n") + "\n" + row
    log_p.write_text(text, encoding="utf-8", newline="\n")
    print("تم التحديث.")


if __name__ == "__main__":
    main()
